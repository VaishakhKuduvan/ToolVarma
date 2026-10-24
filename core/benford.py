"""
Benford's Law engine (CAG of India research paper — Nigrini 5-test + SA 320 tiers).
Pure logic: digit extraction, statistics, Plotly charts and the Excel workpaper.
Moved unchanged from BENFORD_MATERIALITY_NEWCHARTS_ANAMOLIES_NOREPETITION.py.
"""

import math
import os
import shutil
from io import BytesIO

import numpy as np
import pandas as pd
import plotly.graph_objects as go
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from plotly.subplots import make_subplots

# On Streamlit Community Cloud, Chromium comes from packages.txt; tell Kaleido where it is.
if "BROWSER_PATH" not in os.environ and shutil.which("chromium"):
    os.environ["BROWSER_PATH"] = shutil.which("chromium")

# ═══════════════════════════════════════════════════════════════════════════════
# CONFIGURATION
# ═══════════════════════════════════════════════════════════════════════════════
REQUIRED_COLS = [
    'Day', 'Date', 'Vch Type', 'Vch No.', 'Component',
    'Sch III Group', 'Sch III Sub Group', 'tally_main_group',
    'tally_sub_group', 'Particulars', 'Narrations',
    'Debit Amount', 'Credit Amount'
]

# Modern colour palette
COLOR_OBSERVED       = '#3B82F6'   # vibrant blue
COLOR_OBSERVED_LIGHT = '#93C5FD'   # light blue
COLOR_EXPECTED       = '#EF4444'   # red
COLOR_SPIKE          = '#F59E0B'   # amber
COLOR_SPIKE_DARK     = '#D97706'   # dark amber
COLOR_GRID           = '#E5E7EB'
COLOR_BG             = '#FFFFFF'
COLOR_TEXT           = '#1F2937'

# ═══════════════════════════════════════════════════════════════════════════════
# BENFORD EXPECTED PROBABILITIES
# ═══════════════════════════════════════════════════════════════════════════════
def benford_first_digit():
    return {d: math.log10(1 + 1/d) for d in range(1, 10)}

def benford_second_digit():
    return {d: sum(math.log10(1 + 1/(10*k + d)) for k in range(1, 10))
            for d in range(0, 10)}

def benford_first_two():
    return {n: math.log10(1 + 1/n) for n in range(10, 100)}

def benford_first_three():
    return {n: math.log10(1 + 1/n) for n in range(100, 1000)}

def benford_last_two():
    return {n: 0.01 for n in range(0, 100)}

# ═══════════════════════════════════════════════════════════════════════════════
# VECTORISED DIGIT EXTRACTION
# ═══════════════════════════════════════════════════════════════════════════════
def extract_digits_vectorized(amounts: pd.Series) -> pd.DataFrame:
    a = amounts.abs().astype(float)
    result = pd.DataFrame(index=amounts.index)

    valid = a > 0
    a_valid = a.where(valid)

    with np.errstate(divide='ignore', invalid='ignore'):
        log10_vals = np.log10(a_valid)

    floor_log = np.floor(log10_vals)
    leading = a_valid / (10.0 ** floor_log)
    leading = leading.where(leading < 10.0, 9.9999999999)

    first       = np.floor(leading)
    first_two   = np.floor(leading * 10)
    first_three = np.floor(leading * 100)
    second      = first_two % 10

    result['first']       = first.astype('Int64')
    result['second']      = second.astype('Int64')
    result['first_two']   = first_two.astype('Int64')
    result['first_three'] = first_three.astype('Int64')

    int_part = a.fillna(0).astype('int64')
    last_two = int_part % 100
    last_two_valid = int_part >= 10
    result['last_two'] = last_two.where(last_two_valid).astype('Int64')

    return result

TEST_TO_COL = {
    '1. First Digit Test'        : 'first',
    '2. Second Digit Test'       : 'second',
    '3. First-Two Digits Test'   : 'first_two',
    '4. First-Three Digits Test' : 'first_three',
    '5. Last-Two Digits Test'    : 'last_two',
}

# ═══════════════════════════════════════════════════════════════════════════════
# STATISTICAL CALCULATIONS
# ═══════════════════════════════════════════════════════════════════════════════
def _gamma_q(a, x):
    """Regularised upper incomplete gamma Q(a, x) (Numerical Recipes series / continued fraction)."""
    if x <= 0:
        return 1.0
    gln = math.lgamma(a)
    if x < a + 1:  # series for P(a, x)
        term = total = 1.0 / a
        ap = a
        for _ in range(1000):
            ap += 1
            term *= x / ap
            total += term
            if abs(term) < abs(total) * 1e-15:
                break
        return max(0.0, 1.0 - total * math.exp(-x + a * math.log(x) - gln))
    b = x + 1 - a  # Lentz continued fraction for Q(a, x)
    c, d = 1e300, 1.0 / b
    h = d
    for i in range(1, 1000):
        an = -i * (i - a)
        b += 2
        d = an * d + b
        d = 1e-300 if abs(d) < 1e-300 else d
        c = b + an / c
        c = 1e-300 if abs(c) < 1e-300 else c
        d = 1.0 / d
        delta = d * c
        h *= delta
        if abs(delta - 1) < 1e-15:
            break
    return math.exp(-x + a * math.log(x) - gln) * h


def chisquare(observed, expected):
    """Pearson chi-square goodness-of-fit, same result as scipy.stats.chisquare (dof = k - 1)."""
    chi2 = float(sum((o - e) ** 2 / e for o, e in zip(observed, expected) if e > 0))
    dof = len(observed) - 1
    return chi2, _gamma_q(dof / 2.0, chi2 / 2.0)


def calc_z(actual_p, expected_p, n):
    if n == 0:
        return 0
    se = math.sqrt(expected_p * (1 - expected_p) / n)
    if se == 0:
        return 0
    diff = abs(actual_p - expected_p)
    correction = 1 / (2 * n)
    if diff > correction:
        diff -= correction
    return diff / se

def calc_mad(observed_props, expected_props):
    diffs = [abs(observed_props.get(k, 0) - v) for k, v in expected_props.items()]
    return sum(diffs) / len(diffs)

MAD_THRESHOLDS = {
    'first':       (0.006,  0.012,  0.015),
    'second':      (0.008,  0.010,  0.012),
    'first_two':   (0.0012, 0.0018, 0.0022),
    'first_three': (0.00036, 0.00044, 0.00050),
    'last_two':    (0.0012, 0.0018, 0.0022),
}

def mad_conformity(mad, test_key):
    t = MAD_THRESHOLDS.get(test_key, (0.006, 0.012, 0.015))
    if mad < t[0]:   return "Close Conformity"
    elif mad < t[1]: return "Acceptable Conformity"
    elif mad < t[2]: return "Marginally Acceptable"
    else:            return "Nonconformity"

# ═══════════════════════════════════════════════════════════════════════════════
# CORE TEST RUNNER
# ═══════════════════════════════════════════════════════════════════════════════
def run_test_with_digits(digit_series, expected_probs, test_name, test_key, digit_col):
    digits = digit_series.dropna().astype(int)
    n = len(digits)
    if n == 0:
        return None

    counts = digits.value_counts().to_dict()
    obs_counts = [counts.get(k, 0) for k in expected_probs.keys()]
    exp_counts = [expected_probs[k] * n for k in expected_probs.keys()]
    obs_props  = {k: counts.get(k, 0)/n for k in expected_probs.keys()}

    try:
        chi2, p_val = chisquare(obs_counts, exp_counts)
    except Exception:
        chi2, p_val = np.nan, np.nan

    mad = calc_mad(obs_props, expected_probs)
    conformity = mad_conformity(mad, test_key)
    z_stats = {k: calc_z(obs_props[k], expected_probs[k], n)
               for k in expected_probs.keys()}

    df = pd.DataFrame({
        'Digit'         : list(expected_probs.keys()),
        'Observed Count': obs_counts,
        'Expected Count': [round(e, 2) for e in exp_counts],
        'Observed %'    : [obs_props[k]*100 for k in expected_probs.keys()],
        'Expected %'    : [expected_probs[k]*100 for k in expected_probs.keys()],
        'Z-Statistic'   : [z_stats[k] for k in expected_probs.keys()],
    })
    df['Flag'] = df['Z-Statistic'].apply(lambda z: 'SPIKE' if z > 1.96 else '')
    df['Deviation %'] = df['Observed %'] - df['Expected %']
    spike_digits = df.loc[df['Z-Statistic'] > 1.96, 'Digit'].tolist()

    return {
        'name': test_name, 'df': df, 'n': n,
        'chi2': chi2, 'p_val': p_val,
        'mad': mad, 'conformity': conformity,
        'spike_digits': spike_digits,
        'digit_col': digit_col,
    }

# ═══════════════════════════════════════════════════════════════════════════════
# MODERN PLOTLY VISUALISATION
# ═══════════════════════════════════════════════════════════════════════════════
def make_plotly_chart(result, tier_label="", show_deviation=True):
    """
    Build a modern, interactive Plotly chart for a Benford test result.
    Returns a plotly.graph_objects.Figure.
    """
    df = result['df']
    digits = df['Digit'].astype(str).tolist()
    n_digits = len(digits)

    # Conditional bar colours: spikes in amber, normal in blue
    bar_colors = [
        COLOR_SPIKE if z > 1.96 else COLOR_OBSERVED
        for z in df['Z-Statistic']
    ]
    border_colors = [
        COLOR_SPIKE_DARK if z > 1.96 else COLOR_OBSERVED
        for z in df['Z-Statistic']
    ]

    # Rich hover text for each bar
    hover_texts = [
        f"<b>Digit: {d}</b><br>"
        f"Observed: {oc:,} ({op:.3f}%)<br>"
        f"Expected: {ec:,.1f} ({ep:.3f}%)<br>"
        f"Deviation: {dev:+.3f}%<br>"
        f"Z-Statistic: {z:.3f}<br>"
        f"<b>{'Spike (Z > 1.96)' if z > 1.96 else 'Within bounds'}</b>"
        for d, oc, op, ec, ep, dev, z in zip(
            df['Digit'], df['Observed Count'], df['Observed %'],
            df['Expected Count'], df['Expected %'], df['Deviation %'],
            df['Z-Statistic']
        )
    ]

    # Two-panel layout: main chart on top, deviation strip on bottom
    if show_deviation:
        fig = make_subplots(
            rows=2, cols=1, shared_xaxes=True,
            row_heights=[0.78, 0.22],
            vertical_spacing=0.06,
            subplot_titles=(None, "Deviation from Benford Expected (%)"),
        )
    else:
        fig = make_subplots(rows=1, cols=1)

    # ── Main: Observed bars ──
    fig.add_trace(
        go.Bar(
            x=digits, y=df['Observed %'],
            name='Observed %',
            marker=dict(
                color=bar_colors,
                line=dict(color=border_colors, width=1.2),
                opacity=0.92,
            ),
            hovertext=hover_texts,
            hoverinfo='text',
            showlegend=True,
        ),
        row=1, col=1,
    )

    # ── Main: Expected line (Benford) ──
    fig.add_trace(
        go.Scatter(
            x=digits, y=df['Expected %'],
            name='Benford Expected %',
            mode='lines+markers',
            line=dict(color=COLOR_EXPECTED, width=2.5, shape='spline', smoothing=0.6),
            marker=dict(size=7, color=COLOR_EXPECTED,
                        line=dict(color='white', width=1.5)),
            hovertemplate='<b>Digit: %{x}</b><br>Expected: %{y:.3f}%<extra></extra>',
        ),
        row=1, col=1,
    )

    # ── Main: Highlight spike points with star markers ──
    spike_df = df[df['Z-Statistic'] > 1.96]
    if len(spike_df):
        fig.add_trace(
            go.Scatter(
                x=spike_df['Digit'].astype(str), y=spike_df['Observed %'],
                name='Spike (Z > 1.96)',
                mode='markers',
                marker=dict(
                    symbol='star', size=16,
                    color=COLOR_SPIKE,
                    line=dict(color=COLOR_SPIKE_DARK, width=2),
                ),
                hovertemplate='<b>SPIKE Digit %{x}</b><br>Obs %{y:.3f}%<extra></extra>',
            ),
            row=1, col=1,
        )

    # ── Bottom: Deviation bars (red below 0, green above) ──
    if show_deviation:
        dev_colors = [
            ('#10B981' if d >= 0 else '#EF4444') for d in df['Deviation %']
        ]
        fig.add_trace(
            go.Bar(
                x=digits, y=df['Deviation %'],
                name='Deviation %',
                marker=dict(color=dev_colors, opacity=0.7,
                            line=dict(color='rgba(0,0,0,0.2)', width=0.5)),
                hovertemplate='<b>Digit %{x}</b><br>Deviation: %{y:+.3f}%<extra></extra>',
                showlegend=False,
            ),
            row=2, col=1,
        )
        fig.add_hline(y=0, line_color='gray', line_width=1, row=2, col=1)

    # ── Title with KPI badges ──
    conf_color = {
        'Close Conformity':       '#10B981',
        'Acceptable Conformity':  '#3B82F6',
        'Marginally Acceptable':  '#F59E0B',
        'Nonconformity':          '#EF4444',
    }.get(result['conformity'], '#6B7280')

    tier_clause = f"  •  Tier: <b>{tier_label}</b>" if tier_label else ""
    main_title = (
        f"<b style='font-size:16px'>{result['name']}</b>{tier_clause}<br>"
        f"<span style='font-size:12px; color:{COLOR_TEXT}'>"
        f"n = <b>{result['n']:,}</b>  •  "
        f"MAD = <b>{result['mad']:.5f}</b>  •  "
        f"χ² = <b>{result['chi2']:.2f}</b>  •  "
        f"p = <b>{result['p_val']:.4f}</b>  •  "
        f"<span style='color:{conf_color}'>● {result['conformity']}</span>  •  "
        f"Spikes = <b>{len(result['spike_digits'])}</b>"
        f"</span>"
    )

    # Determine sensible tick spacing for large digit ranges
    tickangle = -45 if n_digits > 20 else 0
    if n_digits > 200:      # first-three digits
        dtick = 50
    elif n_digits > 50:     # first-two / last-two
        dtick = 5
    else:
        dtick = 1

    fig.update_layout(
        title=dict(text=main_title, x=0.01, xanchor='left',
                   font=dict(family='Inter, Segoe UI, sans-serif', size=14)),
        height=620 if show_deviation else 460,
        plot_bgcolor=COLOR_BG,
        paper_bgcolor=COLOR_BG,
        font=dict(family='Inter, Segoe UI, sans-serif', size=11, color=COLOR_TEXT),
        margin=dict(l=60, r=30, t=110, b=60),
        legend=dict(
            orientation='h', yanchor='bottom', y=1.02,
            xanchor='right', x=1.0,
            bgcolor='rgba(255,255,255,0.85)',
            bordercolor=COLOR_GRID, borderwidth=1,
        ),
        hovermode='x unified',
        bargap=0.15,
    )

    fig.update_xaxes(
        title_text='Digit', row=2 if show_deviation else 1, col=1,
        gridcolor=COLOR_GRID, linecolor=COLOR_GRID,
        tickangle=tickangle, dtick=dtick,
        showspikes=True, spikemode='across',
        spikecolor='#9CA3AF', spikedash='dot', spikethickness=1,
    )
    if show_deviation:
        fig.update_xaxes(showticklabels=False, row=1, col=1,
                         gridcolor=COLOR_GRID, linecolor=COLOR_GRID)

    fig.update_yaxes(
        title_text='Frequency (%)',
        gridcolor=COLOR_GRID, linecolor=COLOR_GRID,
        zeroline=True, zerolinecolor=COLOR_GRID,
        row=1, col=1,
    )
    if show_deviation:
        fig.update_yaxes(
            title_text='Deviation (%)',
            gridcolor=COLOR_GRID, linecolor=COLOR_GRID,
            row=2, col=1,
        )

    return fig

def for_screen(fig):
    """Adapt a chart built for the Excel export (white, titled) to the app, where the test name and
    statistics are shown above it and the background must follow the light / dark theme."""
    grid = 'rgba(128, 128, 128, 0.22)'
    fig.update_layout(
        title=None, margin=dict(l=56, r=16, t=48, b=48),
        paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)',
        font=dict(family='Inter, Segoe UI, sans-serif', color=None),
        legend=dict(bgcolor='rgba(0,0,0,0)', borderwidth=0),
    )
    fig.update_xaxes(gridcolor=grid, linecolor=grid)
    fig.update_yaxes(gridcolor=grid, linecolor=grid, zerolinecolor=grid)
    fig.update_annotations(font=dict(size=12))
    return fig


def plotly_to_png_bytes(fig, width=1200, height=600, scale=2):
    """Export a Plotly figure to PNG bytes for Excel embedding via Kaleido."""
    buf = BytesIO()
    img_bytes = fig.to_image(format='png', width=width, height=height, scale=scale)
    buf.write(img_bytes)
    buf.seek(0)
    return buf

# ═══════════════════════════════════════════════════════════════════════════════
# NARRATIVE
# ═══════════════════════════════════════════════════════════════════════════════
def summary_paragraph(result, col_name, tier_label, tier_threshold):
    mad = result['mad']; conf = result['conformity']; p = result['p_val']
    name = result['name']
    spikes = result['df'][result['df']['Z-Statistic'] > 1.96]
    spike_digits = ', '.join(map(str, spikes['Digit'].head(8).tolist()))

    tier_clause = (f" applied on the **{tier_label}** tier (≥ ₹{tier_threshold:,.2f})"
                   if tier_threshold > 0 else " applied on the **Full Population**")

    line1 = (f"The **{name}**{tier_clause} on `{col_name}` with **n = {result['n']:,}** "
             f"records yielded MAD = **{mad:.5f}** indicating **{conf}** "
             f"(Chi-square p-value = {p:.4f}).")
    if len(spikes) > 0:
        line2 = (f"**{len(spikes)} digit-bucket(s)** show significant deviation "
                 f"(Z > 1.96): **{spike_digits}**.")
        line3 = ("Since these spikes survive the materiality filter, they represent "
                 "**audit-relevant anomalies** — examine underlying vouchers for split "
                 "transactions, threshold-evasion, duplicates or fabricated entries.")
    else:
        line2 = "No digit-bucket crosses the Z > 1.96 critical value at 5% significance."
        line3 = ("The dataset broadly conforms to Benford expectations at this materiality "
                 "tier; however, conformity is not a guarantee of absence of fraud.")
    return f"{line1}  \n{line2}  \n{line3}"

# ═══════════════════════════════════════════════════════════════════════════════
# DE-DUPLICATION
# ═══════════════════════════════════════════════════════════════════════════════
TIER_PRIORITY = {
    'Above OM': 5, 'Above SM': 4, 'Above PM': 3,
    'Above CTT': 2, 'Full Population': 1,
}
TEST_PRIORITY = {
    '4. First-Three Digits Test': 5,
    '3. First-Two Digits Test'  : 4,
    '5. Last-Two Digits Test'   : 3,
    '2. Second Digit Test'      : 2,
    '1. First Digit Test'       : 1,
}

def collect_all_flags(tier_outputs):
    events = []
    for tier_label, payload in tier_outputs.items():
        src_df    = payload['source_df']
        digits_df = payload['digits_df']
        if len(src_df) == 0:
            continue
        for r in payload['results']:
            spike_digits = r['spike_digits']
            if not spike_digits:
                continue
            digit_col = r['digit_col']
            digits = digits_df[digit_col]
            mask = digits.isin(spike_digits)
            flagged_idx = src_df.index[mask]
            flagged_digits = digits.loc[flagged_idx]
            for idx, d in zip(flagged_idx, flagged_digits):
                events.append({
                    'row_idx'     : idx,
                    'tier'        : tier_label,
                    'test'        : r['name'],
                    'digit_bucket': d,
                })
    return pd.DataFrame(events)

def resolve_winning_placement(flags_df, last_two_orthogonal=True):
    if flags_df.empty:
        return flags_df.assign(winning=False)
    flags_df = flags_df.copy()
    flags_df['tier_rank'] = flags_df['tier'].map(TIER_PRIORITY).fillna(0)
    flags_df['test_rank'] = flags_df['test'].map(TEST_PRIORITY).fillna(0)
    if last_two_orthogonal:
        flags_df['track'] = flags_df['test'].apply(
            lambda t: 'last_two' if t == '5. Last-Two Digits Test' else 'leading'
        )
        flags_df.sort_values(
            ['row_idx', 'track', 'tier_rank', 'test_rank'],
            ascending=[True, True, False, False], inplace=True
        )
        winning_keys = flags_df.drop_duplicates(
            subset=['row_idx', 'track'], keep='first'
        ).index
    else:
        flags_df.sort_values(
            ['row_idx', 'tier_rank', 'test_rank'],
            ascending=[True, False, False], inplace=True
        )
        winning_keys = flags_df.drop_duplicates(subset='row_idx', keep='first').index
    flags_df['winning'] = flags_df.index.isin(winning_keys)
    return flags_df

# ═══════════════════════════════════════════════════════════════════════════════
# EXCEL EXPORT
# ═══════════════════════════════════════════════════════════════════════════════
def safe_sheet(name):
    for ch in r'[]:*?/\\':
        name = name.replace(ch, '')
    return name[:31]

def build_excel_report(tier_outputs, materiality, target_col, vch_filter,
                       original_cols, dedup_mode='strict', include_charts=True,
                       chart_buffers=None):
    output = BytesIO()

    flags_df = collect_all_flags(tier_outputs)
    if dedup_mode == 'none':
        if not flags_df.empty:
            flags_df = flags_df.assign(winning=True)
    elif dedup_mode == 'split':
        flags_df = resolve_winning_placement(flags_df, last_two_orthogonal=True)
    else:
        flags_df = resolve_winning_placement(flags_df, last_two_orthogonal=False)

    with pd.ExcelWriter(output, engine='openpyxl') as writer:

        # ────────────── Materiality_Setup ──────────────
        meta_df = pd.DataFrame({
            'Parameter': ['Report Title', 'Target Column', 'Vch Type Filter',
                          'Critical Z-value', 'Significance Level', 'Framework',
                          'De-duplication Mode', 'Charts Embedded'],
            'Value':     ["Benford's Law Audit Analysis (CAG + SA 320)",
                          target_col, vch_filter, '1.96', '5%',
                          'Nigrini 5-Test + Materiality Tiers',
                          dedup_mode.upper(), 'YES' if include_charts else 'NO'],
        })
        mat_df = pd.DataFrame({
            'Materiality Level': ['Overall Materiality (OM)',
                                  'Performance Materiality (PM)',
                                  'Specific Materiality (SM)',
                                  'Clearly Trivial Threshold (CTT)'],
            'Amount (Rs)': [materiality['OM'], materiality['PM'],
                            materiality['SM'], materiality['CTT']],
            'Purpose': [
                'FS-level threshold (SA 320)',
                'Buffer for undetected misstatements (50-75% of OM)',
                'Sensitive / fraud-prone class threshold',
                'Below which findings need not be aggregated (~5% of OM)',
            ],
        })
        meta_df.to_excel(writer, sheet_name='Materiality_Setup',
                         index=False, startrow=0)
        mat_df.to_excel(writer, sheet_name='Materiality_Setup',
                        index=False, startrow=len(meta_df) + 3)

        # ────────────── Consolidated_Summary ──────────────
        all_summary_rows = []
        for tier_label, payload in tier_outputs.items():
            for r in payload['results']:
                if flags_df.empty:
                    allocated = 0
                else:
                    allocated = int(((flags_df['tier'] == tier_label) &
                                     (flags_df['test'] == r['name']) &
                                     (flags_df['winning'])).sum())
                all_summary_rows.append({
                    'Materiality Tier'           : tier_label,
                    'Threshold (Rs)'             : payload['threshold'],
                    'Test'                       : r['name'],
                    'Records (n)'                : r['n'],
                    'MAD'                        : round(r['mad'], 6),
                    'Chi-square'                 : round(r['chi2'], 4) if not np.isnan(r['chi2']) else 'N/A',
                    'p-value'                    : round(r['p_val'], 6) if not np.isnan(r['p_val']) else 'N/A',
                    'Conformity'                 : r['conformity'],
                    'Spike Buckets'              : len(r['spike_digits']),
                    'Spike Digits'               : ', '.join(map(str, r['spike_digits'])) if r['spike_digits'] else 'None',
                    'Rows Allocated (Post-Dedup)': allocated,
                })
        pd.DataFrame(all_summary_rows).to_excel(
            writer, sheet_name='Consolidated_Summary', index=False)

        # ────────────── Master_Unique_Anomalies ──────────────
        if not flags_df.empty:
            agg_rows = []
            for row_idx, grp in flags_df.groupby('row_idx'):
                winning_row = grp[grp['winning']]
                placed_tier = winning_row['tier'].iloc[0] if len(winning_row) else ''
                placed_test = winning_row['test'].iloc[0] if len(winning_row) else ''
                agg_rows.append({
                    'row_idx'           : row_idx,
                    'Triggered_In_Tiers': ', '.join(sorted(set(grp['tier']))),
                    'Triggered_In_Tests': ', '.join(sorted(set(grp['test']))),
                    'Total_Triggers'    : len(grp),
                    'Placed_In_Tier'    : placed_tier,
                    'Placed_In_Test'    : placed_test,
                })
            agg_df = pd.DataFrame(agg_rows).set_index('row_idx')

            union_df = pd.concat([p['source_df'] for p in tier_outputs.values()])
            union_df = union_df[~union_df.index.duplicated(keep='first')]

            common_idx = union_df.index.intersection(agg_df.index)
            master = union_df.loc[common_idx, original_cols].copy()
            master = master.join(agg_df, how='left')

            master['Above_CTT'] = master[target_col] >= materiality['CTT']
            master['Above_PM']  = master[target_col] >= materiality['PM']
            master['Above_SM']  = master[target_col] >= materiality['SM']
            master['Above_OM']  = master[target_col] >= materiality['OM']

            master.reset_index(drop=True, inplace=True)
            master.to_excel(writer, sheet_name='Master_Unique_Anomalies', index=False)
        else:
            pd.DataFrame({'Note': ['No anomalies detected across any test/tier.']}) \
              .to_excel(writer, sheet_name='Master_Unique_Anomalies', index=False)

        # ────────────── Charts Sheet ──────────────
        if include_charts:
            pd.DataFrame({
                'Info': ['All Benford analysis charts are embedded in this sheet. '
                         'Scroll down to view each Tier × Test combination.']
            }).to_excel(writer, sheet_name='Charts', index=False)
            charts_ws = writer.sheets['Charts']
            current_row = 4
            for tier_label, payload in tier_outputs.items():
                for r in payload['results']:
                    header_cell = charts_ws.cell(
                        row=current_row, column=1,
                        value=f"{tier_label}  |  {r['name']}  "
                              f"|  n={r['n']:,}  |  MAD={r['mad']:.5f}  "
                              f"|  {r['conformity']}")
                    header_cell.font = header_cell.font.copy(bold=True, size=12)

                    img_buf = None
                    if chart_buffers is not None:
                        cached = chart_buffers.get((tier_label, r['name']))
                        if cached is not None:
                            cached.seek(0)
                            img_buf = BytesIO(cached.getvalue())
                    if img_buf is None:
                        fig = make_plotly_chart(r, tier_label)
                        img_buf = plotly_to_png_bytes(fig)

                    try:
                        img = XLImage(img_buf)
                        img.width = 960
                        img.height = 480
                        charts_ws.add_image(img, f"A{current_row + 1}")
                    except Exception:
                        pass
                    current_row += 28

        # ────────────── Per-tier × test spike sheets ──────────────
        TEST_SHORT = {
            '1. First Digit Test'        : '1FD',
            '2. Second Digit Test'       : '2SD',
            '3. First-Two Digits Test'   : '3F2D',
            '4. First-Three Digits Test' : '4F3D',
            '5. Last-Two Digits Test'    : '5L2D',
        }
        TIER_SHORT = {
            'Full Population': 'FULL', 'Above CTT': 'CTT',
            'Above PM': 'PM', 'Above SM': 'SM', 'Above OM': 'OM',
        }

        ref_records = []
        for tier_label, payload in tier_outputs.items():
            src_df    = payload['source_df']
            digits_df = payload['digits_df']
            for r in payload['results']:
                sheet = safe_sheet(
                    f"Sp_{TIER_SHORT.get(tier_label, tier_label[:3])}_"
                    f"{TEST_SHORT.get(r['name'], r['name'][:3])}"
                )

                if flags_df.empty:
                    keep_idx = []
                else:
                    keep_idx = flags_df[
                        (flags_df['tier'] == tier_label) &
                        (flags_df['test'] == r['name']) &
                        (flags_df['winning'])
                    ]['row_idx'].tolist()

                if not keep_idx:
                    pd.DataFrame({
                        'Note': [f"No unique spikes allocated to {r['name']} in {tier_label} "
                                 f"(dedup mode: '{dedup_mode}'). "
                                 f"Anomalies, if any, were placed in a more specific sheet — "
                                 f"see Master_Unique_Anomalies."]
                    }).to_excel(writer, sheet_name=sheet, index=False)
                    continue

                valid_idx = src_df.index.intersection(keep_idx)
                spike_rows = src_df.loc[valid_idx, original_cols].copy()
                spike_rows['Spike_Digit_Bucket'] = digits_df.loc[valid_idx, r['digit_col']]
                spike_rows['Above_CTT'] = spike_rows[target_col] >= materiality['CTT']
                spike_rows['Above_PM']  = spike_rows[target_col] >= materiality['PM']
                spike_rows['Above_SM']  = spike_rows[target_col] >= materiality['SM']
                spike_rows['Above_OM']  = spike_rows[target_col] >= materiality['OM']

                sort_cols = ['Spike_Digit_Bucket']
                if 'Date' in spike_rows.columns:
                    sort_cols.append('Date')
                spike_rows.sort_values(by=sort_cols, inplace=True)
                spike_rows.to_excel(writer, sheet_name=sheet, index=False)

                if include_charts:
                    try:
                        spike_ws = writer.sheets[sheet]
                        img_buf = None
                        if chart_buffers is not None:
                            cached = chart_buffers.get((tier_label, r['name']))
                            if cached is not None:
                                cached.seek(0)
                                img_buf = BytesIO(cached.getvalue())
                        if img_buf is None:
                            fig = make_plotly_chart(r, tier_label)
                            img_buf = plotly_to_png_bytes(fig, width=1000, height=500, scale=1.5)
                        img = XLImage(img_buf)
                        img.width = 800
                        img.height = 400
                        anchor_col = get_column_letter(len(spike_rows.columns) + 3)
                        spike_ws.add_image(img, f"{anchor_col}2")
                    except Exception:
                        pass

                bucket_counts = spike_rows['Spike_Digit_Bucket'].value_counts().to_dict()
                for d, cnt in bucket_counts.items():
                    ref_records.append({
                        'Materiality Tier'  : tier_label,
                        'Test'              : r['name'],
                        'Spike Digit Bucket': d,
                        'Records in Bucket' : cnt,
                        'Sheet Name'        : sheet,
                    })

        # ────────────── Spike_Reference_Index ──────────────
        if ref_records:
            pd.DataFrame(ref_records).to_excel(
                writer, sheet_name='Spike_Reference_Index', index=False)
        else:
            pd.DataFrame({'Note': ['No spikes after de-duplication.']}) \
              .to_excel(writer, sheet_name='Spike_Reference_Index', index=False)

    output.seek(0)
    return output


# ═══════════════════════════════════════════════════════════════════════════════
# TIER RUNNER (used by the Streamlit page)
# ═══════════════════════════════════════════════════════════════════════════════
TESTS_CONFIG = [
    ('1. First Digit Test',        benford_first_digit(),  'first',       'first'),
    ('2. Second Digit Test',       benford_second_digit(), 'second',      'second'),
    ('3. First-Two Digits Test',   benford_first_two(),    'first_two',   'first_two'),
    ('4. First-Three Digits Test', benford_first_three(),  'first_three', 'first_three'),
    ('5. Last-Two Digits Test',    benford_last_two(),     'last_two',    'last_two'),
]
TIER_LABELS = ['Full Population', 'Above CTT', 'Above PM', 'Above SM', 'Above OM']
MIN_TIER_RECORDS = 30


def tier_thresholds(materiality):
    return {
        'Full Population': 0.0,
        'Above CTT': materiality['CTT'],
        'Above PM' : materiality['PM'],
        'Above SM' : materiality['SM'],
        'Above OM' : materiality['OM'],
    }


def run_tiers(df_pos, target_col, thresholds, tier_choice, progress_cb=None):
    """Run all five tests on every selected tier.
    Returns (tier_outputs, skipped); skipped maps tier -> record count when below MIN_TIER_RECORDS."""
    tier_outputs, skipped = {}, {}
    total = max(1, len(tier_choice))
    for ti, tier_label in enumerate(tier_choice):
        threshold = thresholds[tier_label]
        tier_df = df_pos[df_pos[target_col] >= threshold].copy()
        if progress_cb:
            progress_cb((ti + 0.1) / total, f"Processing {tier_label} ({len(tier_df):,} records)...")
        if len(tier_df) < MIN_TIER_RECORDS:
            skipped[tier_label] = len(tier_df)
            continue
        digits_df = extract_digits_vectorized(tier_df[target_col])
        results = []
        for name, probs, digit_col, key in TESTS_CONFIG:
            r = run_test_with_digits(digits_df[digit_col], probs, name, key, digit_col)
            if r is not None:
                results.append(r)
        tier_outputs[tier_label] = {
            'threshold' : threshold,
            'results'   : results,
            'source_df' : tier_df,
            'digits_df' : digits_df,
        }
        if progress_cb:
            progress_cb((ti + 1) / total, f"Completed {tier_label}")
    return tier_outputs, skipped
