"""
Ledger AI engine — Gemini-powered ledger mapping, risk scoring & analysis.
Pure logic moved from ledger_ai_studio.py (the Tkinter GUI is replaced by the Streamlit page).

Expected input layout (columns A to G, header in row 1):
    A Ledger Name | B Opening Dr | C Opening Cr | D Current Dr | E Current Cr
    F Closing Dr  | G Closing Cr
"""

import datetime
import json
import logging
import math
import re
import threading
import time
from pathlib import Path

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #
APP_NAME = "Ledger AI Studio"
APP_VERSION = "1.1"

API_BASE = "https://generativelanguage.googleapis.com/v1beta"
AUTO_MODEL_LABEL = "Auto (recommended)"

# Preferred order for automatic fallback. Any other Gemini model that the key can
# access (discovered through the API) is appended after these, so the chain
# adapts automatically when Google adds or retires models.
PREFERRED_MODELS = [
    "gemini-2.5-flash",
    "gemini-flash-latest",
    "gemini-2.5-flash-lite",
    "gemini-2.5-pro",
    "gemini-2.0-flash",
    "gemini-2.0-flash-lite",
    "gemini-flash-lite-latest",
    "gemini-pro-latest",
]
# Models that are not text-generation chat models (or are unsuitable here)
EXCLUDED_MODEL_WORDS = (
    "embedding", "tts", "image", "imagen", "live", "audio", "aqa", "robotics",
    "computer-use", "learnlm", "gemma", "veo", "vision", "native-audio",
)

# Mapping categories. Remove "EQUITY" here if you want the original 4-way split.
ALLOWED_MAPPINGS = ["ASSET", "LIABILITY", "EQUITY", "INCOME", "EXPENSE", "UNCLASSIFIED"]
MAPPING_SYNONYMS = {
    "ASSETS": "ASSET", "LIABILITIES": "LIABILITY", "EXPENSES": "EXPENSE",
    "REVENUE": "INCOME", "CAPITAL": "EQUITY", "RESERVES": "EQUITY",
    "UNKNOWN": "UNCLASSIFIED",
}

DEFAULT_BATCH_SIZE = 40
MAX_ATTEMPTS_PER_MODEL = 2
MAX_ANALYSIS_CHARS = 700_000
MAX_PREVIEW_ROWS = 5000

RISK_COLORS = {5: "F8CBAD", 4: "FCE4D6", 3: "FFF2CC", 2: "E2EFDA", 1: "C6E0B4"}
RISK_TREE_COLORS = {5: "#f4b6a0", 4: "#fbd9c5", 3: "#fff2cc", 2: "#e2efda", 1: "#c6e0b4"}

LOG_DIR = Path.home() / "LedgerAIStudio"

FATAL_MARKERS = (
    "api key not valid", "api_key_invalid", "api key expired", "api key not found",
    "api key is invalid", "has not been used in project", "service_disabled",
)

# --------------------------------------------------------------------------- #
# Prompts
# --------------------------------------------------------------------------- #
STANDARD_IND_AS = "Schedule III - Division II (Ind AS)"
STANDARD_AS = "Schedule III - Division I (AS)"
STANDARD_BASIC = "Basic (Asset / Liability / Equity / Income / Expense)"
MAPPING_STANDARDS = [STANDARD_IND_AS, STANDARD_AS, STANDARD_BASIC]
S3_COLUMNS = ["Schedule III Component", "Schedule III Main Group", "Schedule III Sub Group"]

STYLE_FINDINGS = "Findings table in input format + summary (recommended)"
STYLE_FREE = "Free-form text answer"
ANALYSIS_STYLES = [STYLE_FINDINGS, STYLE_FREE]

_MAPPING_RULES = ("""You are a senior Chartered Accountant and forensic auditor reviewing a company's trial balance.
For every ledger you receive, decide (a) its Account Mapping and (b) a Risk Score, using professional judgement.

ACCOUNT MAPPING - choose exactly one of: """ + ", ".join(ALLOWED_MAPPINGS) + """
- Decide primarily from the ledger NAME (accounting nature), then use the balances as supporting evidence.
- ASSET: cash, bank, receivables, inventory, fixed assets, investments, loans/advances given, TDS/GST receivable, prepaid, deposits.
- LIABILITY: payables, borrowings, provisions, statutory dues, GST/TDS payable, advances received, deferred tax liability.
- EQUITY: share capital, securities premium, reserves & surplus, retained earnings, capital reserve, OCI reserves, and appropriations of profit
  such as dividend paid/declared and transfers to reserves (these are NOT expenses).
- INCOME: sales, revenue, other income, interest/dividend received, gains.
- EXPENSE: purchases, cost of goods sold, employee/admin/finance costs, depreciation, taxes charged to P&L, losses.
- UNCLASSIFIED: only when the name is genuinely ambiguous AND the balances give no help, or for total/sub-total rows.
- Ledger names are typically from Indian companies (Schedule III / Ind AS / GST / TDS terminology).

RISK SCORE - integer 1 (lowest) to 5 (highest)
- 5 = Critical: opening and closing balances on opposite sides (Dr<->Cr reversal), or a balance clearly abnormal for the nature of the account
     (e.g. credit balance in cash/bank/inventory, debit balance in share capital or sundry creditors).
- 4 = High: current-period movement more than 2x the opening balance, large unexplained activity, suspense/clearing/miscellaneous accounts with balances,
     related-party or round-sum anomalies.
- 3 = Medium: closing balance more than 1.5x the opening balance, or noticeably unusual growth.
- 2 = Low-Medium: current-period movement below 25% of opening balance (dormant / stagnant but with a balance).
- 1 = Low: stable, expected balance behaviour for the account type.
Use these as a baseline and apply judgement; if several apply, give the highest.
""")

_S3_COMMON = """
SCHEDULE III CLASSIFICATION (Companies Act 2013) - ALSO fill three more fields for every ledger:
- "component": exactly "Balance Sheet" or "Statement of Profit and Loss"
- "main_group": the Schedule III major heading the ledger belongs to
- "sub_group": the Schedule III line item / sub-head it rolls up into, using the exact Schedule III wording
  (e.g. "Other equity - Securities premium", "Financial assets - Trade receivables"). Use "" only for total rows.
"""

_S3_IND_AS = """Use DIVISION II (Ind AS) headings:
Balance Sheet
 Equity: Equity share capital; Other equity (Securities premium, Retained earnings, General reserve, Capital reserve, Capital redemption reserve,
   Other reserves, Other comprehensive income); Instruments entirely equity in nature; Money received against share warrants
 Non-current liabilities: Financial liabilities (Borrowings; Lease liabilities; Other financial liabilities); Provisions; Deferred tax liabilities (net); Other non-current liabilities
 Current liabilities: Financial liabilities (Borrowings; Lease liabilities; Trade payables; Other financial liabilities); Other current liabilities; Provisions; Current tax liabilities (net)
 Non-current assets: Property, plant and equipment; Capital work-in-progress; Investment property; Goodwill; Other intangible assets; Intangible assets under development;
   Financial assets (Investments; Trade receivables; Loans; Other financial assets); Deferred tax assets (net); Other non-current assets
 Current assets: Inventories; Financial assets (Investments; Trade receivables; Cash and cash equivalents; Bank balances other than cash and cash equivalents; Loans; Other financial assets);
   Current tax assets (net); Other current assets; Assets held for sale
Statement of Profit and Loss
 Income: Revenue from operations; Other income
 Expenses: Cost of materials consumed; Purchases of stock-in-trade; Changes in inventories of finished goods, work-in-progress and stock-in-trade; Employee benefits expense;
   Finance costs; Depreciation and amortisation expense; Other expenses
 Exceptional items; Tax expense (Current tax; Deferred tax); Other comprehensive income
Dividend paid/declared and transfers to reserves are appropriations within Other equity - Retained earnings (Statement of Changes in Equity), not P&L expenses.
"""

_S3_AS = """Use DIVISION I (Accounting Standards) headings:
Balance Sheet
 Shareholders' funds: Share capital; Reserves and surplus; Money received against share warrants
 Share application money pending allotment
 Non-current liabilities: Long-term borrowings; Deferred tax liabilities (net); Other long-term liabilities; Long-term provisions
 Current liabilities: Short-term borrowings; Trade payables; Other current liabilities; Short-term provisions
 Non-current assets: Property, plant and equipment (Tangible assets; Intangible assets; Capital work-in-progress; Intangible assets under development);
   Non-current investments; Deferred tax assets (net); Long-term loans and advances; Other non-current assets
 Current assets: Current investments; Inventories; Trade receivables; Cash and cash equivalents; Short-term loans and advances; Other current assets
Statement of Profit and Loss
 Income: Revenue from operations; Other income
 Expenses: Cost of materials consumed; Purchases of stock-in-trade; Changes in inventories of finished goods, work-in-progress and stock-in-trade; Employee benefit expense;
   Finance costs; Depreciation and amortization expense; Other expenses
 Exceptional items; Extraordinary items; Tax expense (Current tax; Deferred tax)
Dividend paid/declared and transfers to reserves are appropriations of profit within Reserves and surplus, not P&L expenses.
"""


def build_mapping_system(standard):
    parts = [_MAPPING_RULES]
    s3 = standard in (STANDARD_IND_AS, STANDARD_AS)
    if s3:
        parts.append(_S3_COMMON)
        parts.append(_S3_IND_AS if standard == STANDARD_IND_AS else _S3_AS)
    shape = ('{"results":[{"id":<id from input>,"account_mapping":"<one allowed value>","risk_score":<1-5>,'
             '"rationale":"<max 20 words>"'
             + (',"component":"<...>","main_group":"<...>","sub_group":"<...>"' if s3 else "") + "}]}")
    parts.append("\nOUTPUT - respond with JSON ONLY (no markdown, no commentary), exactly this shape:\n" + shape
                 + "\nReturn one object for EVERY input id, using the same id values. Plain single-line text values only.")
    return "\n".join(parts)


ANALYSIS_SYSTEM = """You are a senior Chartered Accountant and audit professional.
Answer the user's request using ONLY the ledger data supplied. Do not invent figures; quote amounts exactly as given.
Be structured and concise. Use short headings and, where useful, COMPACT markdown tables: a single space around cell text,
a plain |---|---| separator row, and NEVER pad cells with extra spaces or align columns.
Amounts are in the currency of the source file. Columns, when present: Opening / Current (movement) / Closing Debit & Credit,
plus 'Risk Score' (1-5), 'Account Mapping' and Schedule III columns assigned earlier. If asked about Schedule III, refer to
Schedule III of the Companies Act 2013 (Division II = Ind AS, Division I = AS)."""

ANALYSIS_FINDINGS_SYSTEM = """You are a senior Chartered Accountant and audit professional.
Carry out the USER REQUEST using ONLY the ledger data supplied (CSV; the first column 'id' identifies each ledger). Never invent figures.
Respond with JSON ONLY, exactly this shape:
{"overview":"<2-3 sentences>","findings":[{"id":<id>,"finding_type":"<short label e.g. Side reversal / Abnormal movement / Abnormal balance / Dormant balance / Classification query>","details":"<max 40 words, quote amounts exactly as in the data>","suggested_audit_procedures":"<max 40 words, separate procedures with semicolons>"}],"summary":"<short paragraph>"}
Rules: include every ledger that deserves attention (always include those with Risk Score 4 or 5 when present), one entry per id, ids must come from the data;
plain single-line text only - no markdown, no tables, no extra spaces or padding; if no ledger needs attention return an empty findings list;
if the request mentions Schedule III, refer to Schedule III of the Companies Act 2013 (Division II = Ind AS, Division I = AS)."""

DEFAULT_ANALYSIS_PROMPT = (
    "Review this trial balance. Highlight the highest-risk ledgers, abnormal or unexpected balances, opening-vs-closing side reversals and unusual movements. For each high-risk area suggest audit procedures. Present key findings in a table with similar format as displayed in the perview followed by a short summary. Also while perfoming Account mapping refer to Schedule III of Division II of companies act 2013"
)


# --------------------------------------------------------------------------- #
# Exceptions
# --------------------------------------------------------------------------- #
class OperationCancelled(Exception):
    pass


class ApiError(Exception):
    def __init__(self, message, status=None, fatal=False, retriable=False):
        super().__init__(message)
        self.status = status
        self.fatal = fatal          # affects every model (bad key etc.) -> stop immediately
        self.retriable = retriable  # temporary (429/5xx/network) -> retry, then next model


class AllModelsFailed(Exception):
    pass


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def safe_val(val):
    """Convert blanks / text / '1,23,456.50' / '(500)' to float safely (like VBA Val())."""
    try:
        if val is None or (isinstance(val, float) and math.isnan(val)):
            return 0.0
        if isinstance(val, str):
            s = val.strip().replace(",", "").replace("\u20b9", "")
            if not s:
                return 0.0
            neg = s.startswith("(") and s.endswith(")")
            s = s.strip("()")
            num = float(s)
            return -num if neg else num
        if pd.isnull(val):
            return 0.0
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def clean_dataframe(df):
    """Drop empty rows/columns. Raises ValueError if the layout is unusable."""
    df = df.copy()
    df = df.dropna(axis=1, how="all") if any(str(c).startswith("Unnamed") for c in df.columns) else df
    df = df.dropna(how="all").reset_index(drop=True)
    if len(df.columns) < 7:
        raise ValueError("The file must have at least 7 columns (A to G):\n"
                         "Ledger Name, Opening Dr/Cr, Current Dr/Cr, Closing Dr/Cr.")
    keep = []
    for _, row in df.iterrows():
        name_blank = pd.isnull(row.iloc[0]) or str(row.iloc[0]).strip() == ""
        all_zero = all(safe_val(row.iloc[k]) == 0 for k in range(1, 7))
        keep.append(not (name_blank and all_zero))
    df = df[keep].reset_index(drop=True)
    if df.empty:
        raise ValueError("No ledger rows were found in the selected sheet.")
    return df


def read_table(path, sheet=None):
    ext = Path(path).suffix.lower()
    if ext == ".csv":
        return pd.read_csv(path)
    return pd.read_excel(path, sheet_name=sheet if sheet is not None else 0)


def model_sort_key(name):
    if name in PREFERRED_MODELS:
        return (0, PREFERRED_MODELS.index(name), 0, name)
    low = name.lower()
    if "lite" in low:
        tier = 1
    elif "flash" in low:
        tier = 0
    elif "pro" in low:
        tier = 2
    else:
        tier = 3
    m = re.search(r"gemini-(\d+(?:\.\d+)?)", low)
    version = float(m.group(1)) if m else 0.0
    exp_penalty = 1 if ("exp" in low or "preview" in low) else 0
    return (1, tier + exp_penalty * 4, -version, name)


# --------------------------------------------------------------------------- #
# Gemini REST client with automatic model fall-back
# --------------------------------------------------------------------------- #
class GeminiClient:
    def __init__(self, api_key, logger, cancel_event=None):
        self.api_key = api_key.strip()
        self.log = logger
        self.cancel = cancel_event or threading.Event()
        self.session = requests.Session()
        self.session.headers.update({"x-goog-api-key": self.api_key,
                                     "Content-Type": "application/json"})
        self.models = []        # discovered, ranked
        self.sticky = None      # last model that worked -> tried first next time
        self._no_thinking = set()

    # -- low level -------------------------------------------------------- #
    def _sleep(self, seconds):
        end = time.time() + seconds
        while time.time() < end:
            if self.cancel.is_set():
                raise OperationCancelled()
            time.sleep(0.2)

    def _http(self, method, url, **kwargs):
        if self.cancel.is_set():
            raise OperationCancelled()
        try:
            resp = self.session.request(method, url, timeout=(15, 180), **kwargs)
        except requests.exceptions.RequestException as exc:
            raise ApiError(f"network error ({type(exc).__name__})", retriable=True)
        if resp.status_code == 200:
            try:
                return resp.json()
            except ValueError:
                raise ApiError("invalid JSON in server response", retriable=True)
        try:
            msg = resp.json()["error"]["message"]
        except Exception:
            msg = resp.text[:300]
        low = msg.lower()
        fatal = resp.status_code == 401 or any(m in low for m in FATAL_MARKERS)
        retriable = resp.status_code in (429, 500, 502, 503, 504)
        raise ApiError(f"HTTP {resp.status_code}: {msg[:300]}", resp.status_code, fatal, retriable)

    def _with_retries(self, fn, label):
        for attempt in range(1, MAX_ATTEMPTS_PER_MODEL + 1):
            try:
                return fn()
            except ApiError as exc:
                if exc.fatal or not exc.retriable or attempt == MAX_ATTEMPTS_PER_MODEL:
                    raise
                delay = 3 * attempt
                self.log.warning(f"{label}: {exc} - retrying in {delay}s")
                self._sleep(delay)

    # -- models ----------------------------------------------------------- #
    def list_models(self):
        names, token = [], None
        while True:
            params = {"pageSize": 1000}
            if token:
                params["pageToken"] = token
            data = self._with_retries(lambda: self._http("GET", f"{API_BASE}/models", params=params),
                                      "List models")
            for m in data.get("models", []):
                name = m.get("name", "").replace("models/", "")
                methods = m.get("supportedGenerationMethods", [])
                low = name.lower()
                if ("generateContent" in methods and low.startswith("gemini")
                        and not any(w in low for w in EXCLUDED_MODEL_WORDS)):
                    names.append(name)
            token = data.get("nextPageToken")
            if not token:
                break
        self.models = sorted(set(names), key=model_sort_key)
        return self.models

    def model_chain(self, preferred=None):
        if not self.models:
            try:
                self.list_models()
            except ApiError as exc:
                if exc.fatal:
                    raise
                self.log.warning(f"Could not list models ({exc}); using built-in model list")
            if not self.models:
                self.models = list(PREFERRED_MODELS)
        chain = list(self.models)
        first = preferred or self.sticky
        if first:
            if first in chain:
                chain.remove(first)
            chain.insert(0, first)
        return chain

    # -- generation ------------------------------------------------------- #
    def generate(self, prompt, system=None, json_mode=False, preferred=None, validator=None,
                 max_output_tokens=None, thinking_budget=None):
        """Try every available model in turn. Returns (result, model_used)."""
        errors = []
        for model in self.model_chain(preferred):
            if self.cancel.is_set():
                raise OperationCancelled()
            try:
                def call(m=model):
                    return self._generate_once(m, prompt, system, json_mode, max_output_tokens, thinking_budget)
                text = self._with_retries(call, f"Model {model}")
                result = validator(text) if validator else text
                self.sticky = model
                return result, model
            except ApiError as exc:
                if exc.fatal:
                    raise
                errors.append(f"{model}: {exc}")
                self.log.warning(f"Model '{model}' failed ({exc}) - falling back to next model")
            except ValueError as exc:  # validator rejected the output
                errors.append(f"{model}: unusable output ({exc})")
                self.log.warning(f"Model '{model}' returned unusable output ({exc}) - trying next model")
        raise AllModelsFailed("All available Gemini models failed:\n" + "\n".join(errors[-8:]))

    def _generate_once(self, model, prompt, system, json_mode, max_tokens, thinking_budget):
        gen = {"temperature": 0.1}
        if json_mode:
            gen["responseMimeType"] = "application/json"
        if max_tokens:
            gen["maxOutputTokens"] = max_tokens
        use_thinking = (thinking_budget is not None and "flash" in model.lower()
                        and model not in self._no_thinking)
        if use_thinking:
            gen["thinkingConfig"] = {"thinkingBudget": thinking_budget}  # keeps latency low
        body = {"contents": [{"role": "user", "parts": [{"text": prompt}]}], "generationConfig": gen}
        if system:
            body["systemInstruction"] = {"parts": [{"text": system}]}
        url = f"{API_BASE}/models/{model}:generateContent"
        started = time.time()
        try:
            data = self._http("POST", url, json=body)
        except ApiError as exc:
            if use_thinking and exc.status == 400 and "think" in str(exc).lower():
                self.log.warning(f"Model '{model}' does not accept the thinking setting - retrying without it")
                self._no_thinking.add(model)
                gen.pop("thinkingConfig", None)
                data = self._http("POST", url, json=body)
            else:
                raise
        cands = data.get("candidates") or []
        if not cands:
            reason = (data.get("promptFeedback") or {}).get("blockReason", "unknown")
            raise ApiError(f"no response (blocked: {reason})")
        cand = cands[0]
        parts = (cand.get("content") or {}).get("parts") or []
        text = "".join(p.get("text", "") for p in parts if not p.get("thought"))
        if cand.get("finishReason") == "MAX_TOKENS":
            self.log.warning(f"Model '{model}' hit the output limit - the answer may be incomplete")
        self.log.info(f"'{model}' answered in {time.time() - started:.0f}s")
        if not text.strip():
            raise ApiError(f"empty response (finishReason={cand.get('finishReason')})")
        return text


# --------------------------------------------------------------------------- #
# Mapping / analysis engine (no GUI code - easy to test)
# --------------------------------------------------------------------------- #
def load_json_loose(text):
    t = text.strip()
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", t, flags=re.IGNORECASE)
    try:
        return json.loads(t)
    except ValueError:
        starts = [p for p in (t.find("{"), t.find("[")) if p != -1]
        if not starts:
            raise ValueError("no JSON found")
        return json.loads(t[min(starts):max(t.rfind("}"), t.rfind("]")) + 1])


def clean_model_text(text):
    """Remove runaway whitespace padding (a known LLM failure with wide markdown tables)."""
    out = []
    for line in text.splitlines():
        s = line.strip()
        if not s:
            out.append("")
            continue
        indent = " " * min(len(line) - len(line.lstrip(" ")), 8)
        out.append(indent + re.sub(r"[ \t]{2,}", " ", s))
    cleaned = re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip()
    if len(cleaned) < 20:
        raise ValueError("answer is empty after cleaning")
    return cleaned


class LedgerAIEngine:
    def __init__(self, client, logger, preferred=None, standard=STANDARD_BASIC):
        self.client = client
        self.log = logger
        self.preferred = preferred
        self.standard = standard
        self.use_s3 = standard in (STANDARD_IND_AS, STANDARD_AS)
        self.system = build_mapping_system(standard)
        self.models_used = []
        self.cancel = client.cancel

    @staticmethod
    def build_records(df):
        recs = []
        for i in range(len(df)):
            row = df.iloc[i]
            name = "" if pd.isnull(row.iloc[0]) else str(row.iloc[0]).strip()
            v = [round(safe_val(row.iloc[k]), 2) for k in range(1, 7)]
            recs.append({"id": i, "ledger": name, "opening_debit": v[0], "opening_credit": v[1],
                         "current_debit": v[2], "current_credit": v[3],
                         "closing_debit": v[4], "closing_credit": v[5]})
        return recs

    @staticmethod
    def parse_results(text):
        obj = load_json_loose(text)
        if isinstance(obj, dict):
            items = obj.get("results")
            if items is None:
                items = next((v for v in obj.values() if isinstance(v, list)), None)
        else:
            items = obj
        if not isinstance(items, list):
            raise ValueError("JSON has no results list")
        cleaned = {}
        for it in items:
            try:
                rid = int(it["id"])
                risk = int(round(float(it["risk_score"])))
                if not 1 <= risk <= 5:
                    continue
                mapping = str(it.get("account_mapping", "")).strip().upper()
                mapping = MAPPING_SYNONYMS.get(mapping, mapping)
                if mapping not in ALLOWED_MAPPINGS:
                    mapping = "UNCLASSIFIED"
                cleaned[rid] = (mapping, risk, str(it.get("rationale", "")).strip(),
                                str(it.get("component", "") or "").strip(),
                                str(it.get("main_group", "") or "").strip(),
                                str(it.get("sub_group", "") or "").strip())
            except (KeyError, TypeError, ValueError):
                continue
        if not cleaned:
            raise ValueError("no valid result rows")
        return cleaned

    def _map_chunk(self, records, depth=0):
        out = {}
        prompt = ("Classify and risk-score the following ledgers. Input records (JSON):\n"
                  + json.dumps(records, ensure_ascii=False))
        try:
            parsed, model = self.client.generate(prompt, self.system, json_mode=True,
                                                 preferred=self.preferred, validator=self.parse_results,
                                                 max_output_tokens=16384, thinking_budget=512)
        except AllModelsFailed as exc:
            if len(records) > 1 and depth < 2:
                mid = len(records) // 2
                self.log.warning(f"Batch of {len(records)} failed on all models - splitting and retrying")
                out.update(self._map_chunk(records[:mid], depth + 1))
                out.update(self._map_chunk(records[mid:], depth + 1))
            else:
                self.log.error(f"Could not map {len(records)} ledger(s): {exc}")
            return out
        if model not in self.models_used:
            self.models_used.append(model)
        ids = {r["id"] for r in records}
        out.update({k: v for k, v in parsed.items() if k in ids})
        missing = [r for r in records if r["id"] not in out]
        if missing and len(missing) < len(records) and depth < 3:
            self.log.warning(f"AI skipped {len(missing)} ledger(s) - re-requesting them")
            out.update(self._map_chunk(missing, depth + 1))
        return out

    def map_all(self, df, batch_size, progress_cb):
        records = self.build_records(df)
        total = len(records)
        results, fails = {}, 0
        self.log.info(f"Mapping standard: {self.standard}")
        for start in range(0, total, batch_size):
            if self.cancel.is_set():
                raise OperationCancelled()
            chunk = records[start:start + batch_size]
            progress_cb(start, total, f"Mapping ledgers {start + 1}-{start + len(chunk)} of {total}")
            self.log.info(f"Sending ledgers {start + 1}-{start + len(chunk)} of {total} to Gemini")
            got = self._map_chunk(chunk)
            results.update(got)
            fails = fails + 1 if not got else 0
            if fails >= 3:
                raise RuntimeError("Three consecutive batches failed. Check your connection, "
                                   "API key quota and the Logs tab.")
        progress_cb(total, total, "Mapping complete")

        out = df.copy()
        out["Risk Score"] = pd.array([results[i][1] if i in results else pd.NA for i in range(total)],
                                     dtype="Int64")
        out["Account Mapping"] = [results[i][0] if i in results else "UNMAPPED" for i in range(total)]
        if self.use_s3:
            for k, col in enumerate(S3_COLUMNS):
                out[col] = [results[i][3 + k] if i in results else "" for i in range(total)]
        out["AI Rationale"] = [results[i][2] if i in results else "AI did not return a result" for i in range(total)]
        unmapped = sum(1 for i in range(total) if i not in results)
        if unmapped:
            self.log.warning(f"{unmapped} ledger(s) could not be mapped (marked UNMAPPED)")
        return out

    # ----------------------------- analysis -------------------------------- #
    @staticmethod
    def parse_findings(text):
        obj = load_json_loose(text)
        if not isinstance(obj, dict) or not isinstance(obj.get("findings"), list):
            raise ValueError("JSON has no findings list")
        items = []
        for it in obj["findings"]:
            try:
                items.append((int(it["id"]),
                              re.sub(r"\s+", " ", str(it.get("finding_type", ""))).strip(),
                              re.sub(r"\s+", " ", str(it.get("details", ""))).strip(),
                              re.sub(r"\s+", " ", str(it.get("suggested_audit_procedures", ""))).strip()))
            except (KeyError, TypeError, ValueError):
                continue
        return {"overview": re.sub(r"[ \t]{2,}", " ", str(obj.get("overview", ""))).strip(),
                "summary": re.sub(r"[ \t]{2,}", " ", str(obj.get("summary", ""))).strip(),
                "findings": items}

    @staticmethod
    def build_findings_df(df, findings):
        extra = ["Finding Type", "Details", "Suggested Audit Procedures"]
        rows, seen = [], set()
        for fid, ftype, details, procs in findings:
            if fid in seen or not 0 <= fid < len(df):
                continue
            seen.add(fid)
            row = df.iloc[fid].to_dict()
            row.update({"Finding Type": ftype, "Details": details, "Suggested Audit Procedures": procs})
            rows.append(row)
        out = pd.DataFrame(rows, columns=list(df.columns) + extra)
        if "Risk Score" in out.columns and len(out):
            key = pd.to_numeric(out["Risk Score"], errors="coerce").fillna(0).to_numpy()
            out = out.iloc[(-key).argsort(kind="stable")].reset_index(drop=True)
            out["Risk Score"] = pd.to_numeric(out["Risk Score"], errors="coerce").astype("Int64")
        return out

    def analyse(self, user_prompt, df, style):
        """Returns (display_text, model, findings_df_or_None)."""
        data = df.copy()
        data.insert(0, "id", range(len(data)))
        csv = data.to_csv(index=False)
        if len(csv) > MAX_ANALYSIS_CHARS:
            self.log.warning("Data is very large - truncating to the first part for analysis")
            csv = csv[:MAX_ANALYSIS_CHARS].rsplit("\n", 1)[0]
        prompt = f"USER REQUEST:\n{user_prompt}\n\nLedger data (CSV):\n{csv}"
        if style == STYLE_FINDINGS:
            parsed, model = self.client.generate(prompt, ANALYSIS_FINDINGS_SYSTEM, json_mode=True,
                                                 preferred=self.preferred, validator=self.parse_findings,
                                                 max_output_tokens=16384, thinking_budget=1024)
            fdf = self.build_findings_df(df, parsed["findings"])
            text = (f"OVERVIEW\n{parsed['overview']}\n\nSUMMARY\n{parsed['summary']}\n\n"
                    f"{len(fdf)} ledger(s) with findings - see the Findings tab / Findings sheet.")
            return text, model, fdf
        text, model = self.client.generate(prompt, ANALYSIS_SYSTEM, preferred=self.preferred,
                                           validator=clean_model_text, max_output_tokens=16384,
                                           thinking_budget=1024)
        return text, model, None


# --------------------------------------------------------------------------- #
# Excel export
# --------------------------------------------------------------------------- #
def _cell_value(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(v, "item"):
        try:
            return v.item()
        except Exception:
            pass
    return v


TEXT_COL_WIDTHS = {"AI Rationale": 55, "Finding Type": 26, "Details": 60, "Suggested Audit Procedures": 60,
                   "Schedule III Component": 24, "Schedule III Main Group": 28, "Schedule III Sub Group": 38}


def _write_table_sheet(wb, title, df, border, head_fill):
    ws = wb.create_sheet(title)
    headers = [str(c) for c in df.columns]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    risk_idx = headers.index("Risk Score") if "Risk Score" in headers else None
    wrap_idx = {i for i, h in enumerate(headers) if h in TEXT_COL_WIDTHS}
    for r in range(len(df)):
        ws.append([_cell_value(df.iat[r, c]) for c in range(len(headers))])
        for c in ws[r + 2]:
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=(c.column - 1) in wrap_idx)
            if 2 <= c.column <= 7:
                c.number_format = "#,##0.00"
        if risk_idx is not None:
            rv = _cell_value(df.iat[r, risk_idx])
            if rv in RISK_COLORS:
                fill = PatternFill("solid", fgColor=RISK_COLORS[rv])
                ws.cell(r + 2, risk_idx + 1).fill = fill
                if "Account Mapping" in headers:
                    ws.cell(r + 2, headers.index("Account Mapping") + 1).fill = fill
    for i, h in enumerate(headers, start=1):
        if h in TEXT_COL_WIDTHS:
            ws.column_dimensions[get_column_letter(i)].width = TEXT_COL_WIDTHS[h]
        else:
            width = max([len(h)] + [len(str(_cell_value(v) or "")) for v in df.iloc[:500, i - 1]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 12), 45)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    return ws


def export_workbook(path, mapped_df, analysis_text, prompt_text="", analysis_model="", mapping_models=None,
                    findings_df=None):
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E79")

    if mapped_df is not None:
        _write_table_sheet(wb, "Mapped Ledger", mapped_df, border, head_fill)
    if findings_df is not None:
        _write_table_sheet(wb, "Findings", findings_df, border, head_fill)

    if analysis_text:
        ws = wb.create_sheet("AI Analysis")
        ws.column_dimensions["A"].width = 38
        for col in "BCDEFGH":
            ws.column_dimensions[col].width = 22
        ws["A1"] = "AI Analysis"
        ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
        meta = [("Generated", datetime.datetime.now().strftime("%d-%b-%Y %H:%M")),
                ("Model", analysis_model or ", ".join(mapping_models or [])),
                ("Prompt", prompt_text)]
        row = 2
        for k, v in meta:
            ws.cell(row, 1, k).font = Font(bold=True)
            ws.cell(row, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            ws.row_dimensions[row].height = max(15, 15 * math.ceil(len(str(v)) / 110))
            row += 1
        row += 1
        in_table = False
        for raw in clean_model_text(analysis_text).splitlines():
            s = raw.strip()
            if not s:
                in_table = False
                row += 1
                continue
            if s.startswith("|") and s.count("|") >= 2:
                cells = [c.strip().replace("**", "").replace("`", "")[:3000] for c in s.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    continue
                for ci, val in enumerate(cells[:8], start=1):
                    cell = ws.cell(row, ci, val)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    cell.border = border
                    if not in_table:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = head_fill
                ws.row_dimensions[row].height = max(15, 15 * math.ceil(max(len(c) for c in cells) / 24))
                in_table = True
                row += 1
                continue
            in_table = False
            heading = s.startswith("#") or (s.isupper() and len(s) < 40)
            text = re.sub(r"^#+\s*", "", s).replace("**", "").replace("`", "")
            text = re.sub(r"^[\*\-]\s+", "\u2022 ", text)[:30000]
            cell = ws.cell(row, 1, text)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if heading or (s.startswith("**") and s.endswith("**")):
                cell.font = Font(bold=True, size=12, color="1F4E79")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            ws.row_dimensions[row].height = max(15, 15 * math.ceil(len(text) / 135))
            row += 1
    wb.save(path)


# --------------------------------------------------------------------------- #
# Logging
# --------------------------------------------------------------------------- #
class SecretFilter(logging.Filter):
    """Make sure the API key can never appear in the log."""
    def __init__(self):
        super().__init__()
        self.secrets = set()

    def filter(self, record):
        msg = record.getMessage()
        for s in self.secrets:
            if s and s in msg:
                msg = msg.replace(s, "***")
        record.msg, record.args = msg, ()
        return True
